"""Find one way to place a figure row that Word AND WPS both put in the middle.

The two disagree: with ``w:jc="center"`` and zero cell margins, Word centres the
table and WPS parks it about 1.8mm to the right — so every multi-picture row
lands off centre for a teacher who opens the file in WPS. This builds the same
picture row several different ways in one document, using the real sizes of
初三第 23 题, so each way can be measured in both programs.
"""

import io
import os
import sys
import tempfile

sys.path.insert(0, r"C:\Users\17321\.minimax\sessions\mvs_fc946281665f4565a3d0ce7af01bd838\workspace\qbj-src\backend")
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "qb_server.settings")
import django

django.setup()

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm
from PIL import Image

from core.library_export import PREVIEW_FIGURE_GAP_PX, PREVIEW_PX_PER_INCH, _figure_paper_size

TEXT_WIDTH_MM = 178
GAP_MM = PREVIEW_FIGURE_GAP_PX * 25.4 / PREVIEW_PX_PER_INCH
# 初三特殊的平行四边形 第 23 题: three diagrams on one row.
SIZES = [(279, 226), (231, 226), (210, 236)]

FOLLOW = ("w:tblLook", "w:tblCaption", "w:tblDescription")


def pictures():
    out = []
    for index, size in enumerate(SIZES):
        buffer = io.BytesIO()
        Image.new("RGB", size, "red" if index % 2 == 0 else "blue").save(buffer, format="PNG")
        out.append({"bytes": buffer.getvalue(), "size": size})
    return out


def zero_cell_margins(props):
    margins = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement("w:" + side)
        edge.set(qn("w:w"), "0")
        edge.set(qn("w:type"), "dxa")
        margins.append(edge)
    props.insert_element_before(margins, *FOLLOW)


def add_row(document, images, *, centred, width_mm, indent_mm=None, pad_mm=None, fixed_layout=False):
    sizes = [_figure_paper_size(image["size"], TEXT_WIDTH_MM) for image in images]
    widths = [item[0] for item in sizes]
    heights = [item[1] for item in sizes]
    columns = [value + GAP_MM for value in widths]
    table = document.add_table(rows=1, cols=len(images))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER if centred else WD_TABLE_ALIGNMENT.LEFT
    for column, value in zip(table.columns, columns):
        column.width = Mm(value)
    props = table._tbl.tblPr

    declared = props.find(qn("w:tblW"))
    if declared is None:
        declared = OxmlElement("w:tblW")
        props.insert_element_before(
            declared,
            "w:jc",
            "w:tblCellSpacing",
            "w:tblInd",
            "w:tblBorders",
            "w:shd",
            "w:tblLayout",
            "w:tblCellMar",
            *FOLLOW,
        )
    declared.set(qn("w:w"), str(int(round(width_mm * 1440 / 25.4))))
    declared.set(qn("w:type"), "dxa")

    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + side)
        border.set(qn("w:val"), "nil")
        borders.append(border)
    props.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", *FOLLOW)

    if fixed_layout:
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        props.insert_element_before(layout, "w:tblCellMar", *FOLLOW)
    if indent_mm is not None:
        indent = OxmlElement("w:tblInd")
        indent.set(qn("w:w"), str(int(round(indent_mm * 1440 / 25.4))))
        indent.set(qn("w:type"), "dxa")
        props.insert_element_before(indent, "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", *FOLLOW)
    if pad_mm is None:
        zero_cell_margins(props)
    else:
        margins = OxmlElement("w:tblCellMar")
        for side in ("top", "left", "bottom", "right"):
            edge = OxmlElement("w:" + side)
            edge.set(qn("w:w"), str(int(round(pad_mm * 1440 / 25.4))))
            edge.set(qn("w:type"), "dxa")
            margins.append(edge)
        props.insert_element_before(margins, *FOLLOW)

    for index, image in enumerate(images):
        cell = table.cell(0, index)
        cell.width = Mm(columns[index])
        paragraph = cell.paragraphs[0]
        paragraph.alignment = 1
        paragraph.add_run().add_picture(io.BytesIO(image["bytes"]), width=Mm(widths[index]), height=Mm(heights[index]))
        paragraph.paragraph_format.keep_together = True
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    return sum(widths) + GAP_MM * (len(widths) - 1)


def build(path):
    document = Document()
    for section in document.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin = section.right_margin = Mm(16)
    images = pictures()
    content = sum(_figure_paper_size(image["size"], TEXT_WIDTH_MM)[0] for image in images) + GAP_MM * (len(images) - 1)
    table_width = content + GAP_MM
    slack = TEXT_WIDTH_MM - table_width

    document.add_paragraph("A centred, no padding (what the export writes today)")
    add_row(document, images, centred=True, width_mm=table_width, pad_mm=0)

    document.add_paragraph("B centred, Word's default padding")
    add_row(document, images, centred=True, width_mm=table_width + 3.8, pad_mm=1.9)

    document.add_paragraph("C pinned to an explicit indent, nothing to re-centre")
    add_row(document, images, centred=False, width_mm=table_width, indent_mm=slack / 2, pad_mm=0)

    document.add_paragraph("D centred, no padding, fixed layout")
    add_row(document, images, centred=True, width_mm=table_width, pad_mm=0, fixed_layout=True)

    document.save(path)
    print(f"content {content:.2f}mm, table {table_width:.2f}mm, slack {slack:.2f}mm")
    print("wrote", path)


if __name__ == "__main__":
    build(os.path.join(tempfile.gettempdir(), "qb-table-variants.docx"))
