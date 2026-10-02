"""Offline, editable DOCX export from exact published versions.

The browser supplies KaTeX MathML and source offsets, never replacement prose
or image URLs. Every field is checked against its stored publication before
conversion. Failures abort the whole download; no raw-LaTeX fallback exists.
"""

from __future__ import annotations

from copy import deepcopy
import hashlib
import io
import re
import unicodedata
import warnings
import zipfile
import xml.sax
from pathlib import Path
from urllib.parse import quote

from django.conf import settings
from django.core.exceptions import RequestDataTooBig
from django.http import HttpResponse, HttpResponseNotAllowed, JsonResponse
from django.views.decorators.csrf import csrf_exempt

from . import features, library
from .library_browse import BrowseError, normalize_ids
from .library_drafts import DraftError, _print_options, _title
from .models import PublishedQuestion

MAX_BODY = 2 * 1024 * 1024
MAX_IMAGE_BYTES = 32 * 1024 * 1024
MAX_TOTAL_IMAGES = 96 * 1024 * 1024
MAX_FIELD_UNITS = 100_000
MAX_MATHML = 80_000
MATH_NS = "http://www.w3.org/1998/Math/MathML"
OMML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
SEGMENT_KEYS = {"type", "start", "end", "latex", "mathml", "display", "displayGroup"}
GROUPS = (("single_choice", "选择题"), ("multiple_choice", "多选题"),
          ("fill_blank", "填空题"), ("true_false", "判断题"), ("free_response", "解答题"))
FIELDS = {"stem", "answer", "analysis", "origin", *(f"options.{key}" for key in "ABCDE")}
ACCENTS = {"⃗": "⃗", "→": "⃗", "←": "⃖", "↔": "⃡", "^": "̂", "~": "̃", "ˉ": "̅",
           "¯": "̅", "˙": "̇", "¨": "̈", "ˇ": "̌", "˘": "̆", "´": "́", "`": "̀"}
MATHML_ELEMENTS = {"math", "semantics", "annotation", "mi", "mn", "mo", "mrow", "mtext", "mspace", "ms",
                   "mstyle", "mpadded", "mphantom", "mfrac", "msqrt", "mroot", "mfenced", "menclose", "msub",
                   "msup", "msubsup", "munder", "mover", "munderover", "mmultiscripts", "mprescripts", "none",
                   "mtable", "mtr", "mtd", "mlabeledtr", "maction"}
LINE_SPACING = 1.35
PARAGRAPH_GAP_PT = 3
QUESTION_GAP_PT = 8


class ExportError(ValueError):
    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def _fail(where, message, status=400):
    raise ExportError(f"{where}：{message}", status)


def _utf16_map(source: str, where: str) -> dict[int, int]:
    """Only complete Unicode scalar boundaries are accepted as JS offsets."""
    if not isinstance(source, str):
        _fail(where, "字段原文应为文字")
    offsets, total = {0: 0}, 0
    for index, character in enumerate(source):
        code = ord(character)
        if 0xD800 <= code <= 0xDFFF or code in {0xFFFE, 0xFFFF} or (code < 32 and character not in "\t\n\r"):
            _fail(where, "文字含不能写入 Word 的控制字符")
        total += 2 if code > 0xFFFF else 1
        offsets[total] = index + 1
    if total > MAX_FIELD_UNITS:
        _fail(where, "文字过长，请减少本次选题")
    return offsets


def _bounds(item, offsets, start, end, where):
    a, b = item.get("start"), item.get("end")
    if type(a) is not int or type(b) is not int or a not in offsets or b not in offsets or a != start or not a < b <= end:
        _fail(where, "排版片段的位置不连续或切开了一个字符，请重新打开组卷")
    return a, b


def _convert_math_root(root):
    """Native Word accents instead of upstream overscript limit glyphs.

    Standard KaTeX distinguishes mover[accent=true] from genuine overscripts.
    Only that explicit shape is adapted; overset and large-operator limits keep
    their normal meaning. Bases recurse to preserve nested accents and roots.
    """
    from lxml import etree
    import mathml2omml
    root = deepcopy(root)
    replacements = {}
    prefix = "QBEXPORTACCENT"
    while prefix in "".join(root.itertext()):
        prefix += "X"
    for node in list(root.iter(f"{{{MATH_NS}}}mover")):
        if root not in node.iterancestors() or node.get("accent") != "true":
            continue
        if len(node) != 2 or node[1].tag != f"{{{MATH_NS}}}mo" or len(node[1]):
            raise ValueError("unsupported accent")
        glyph = node[1].text or ""
        if glyph not in ACCENTS:
            raise ValueError("unsupported accent")
        base = etree.Element(f"{{{MATH_NS}}}math", nsmap={None: MATH_NS})
        base.append(deepcopy(node[0]))
        inner = _convert_math_root(base)
        accent = etree.Element(f"{{{OMML_NS}}}acc")
        props = etree.SubElement(accent, f"{{{OMML_NS}}}accPr")
        character = etree.SubElement(props, f"{{{OMML_NS}}}chr")
        character.set(f"{{{OMML_NS}}}val", ACCENTS[glyph])
        expression = etree.SubElement(accent, f"{{{OMML_NS}}}e")
        for child in inner:
            expression.append(child)
        marker = prefix + str(len(replacements))
        replacements[marker] = accent
        placeholder = etree.Element(f"{{{MATH_NS}}}mi")
        placeholder.text, placeholder.tail = marker, node.tail
        node.getparent().replace(node, placeholder)
    converted = mathml2omml.convert(etree.tostring(root, encoding="unicode"))
    converted = re.sub(r'(<m:groupChrPr><m:chr m:val="[^"]*"/><m:pos m:val="(?:top|bot)"/>)</m:groupChr>',
                       r'\1</m:groupChrPr>', converted)
    parser = etree.XMLParser(resolve_entities=False, load_dtd=False, no_network=True)
    wrapper = etree.fromstring((f'<root xmlns:m="{OMML_NS}">' + converted + '</root>').encode("utf-8"), parser)
    if len(wrapper) != 1:
        raise ValueError("root")
    omml = wrapper[0]
    if omml.tag != f"{{{OMML_NS}}}oMath" or not list(omml):
        raise ValueError("empty")
    seen = set()
    for run in list(omml.iter(f"{{{OMML_NS}}}r")):
        text = run.find(f"{{{OMML_NS}}}t")
        if text is not None and text.text in replacements:
            seen.add(text.text)
            run.getparent().replace(run, replacements[text.text])
    if seen != set(replacements):
        raise ValueError("accent placeholder")
    for radical in omml.findall(f".//{{{OMML_NS}}}rad"):
        if radical.find(f"{{{OMML_NS}}}deg") is None:
            props = etree.Element(f"{{{OMML_NS}}}radPr")
            hidden = etree.SubElement(props, f"{{{OMML_NS}}}degHide")
            hidden.set(f"{{{OMML_NS}}}val", "1")
            radical.insert(0, props)
            radical.insert(1, etree.Element(f"{{{OMML_NS}}}deg"))
        if radical.find(f"{{{OMML_NS}}}e") is None:
            raise ValueError("radical")
    return omml


def _math(segment, where, *, word_math=True):
    """Adapt two narrowly identified upstream OMML bugs, then validate structure.

    mathml2omml 0.0.2 closes groupChrPr with the wrong tag, and omits deg for
    square roots. The latter loses the radicand in a real LibreOffice render.
    These schema repairs are bounded; arbitrary malformed XML is never repaired.
    """
    from lxml import etree

    latex, raw = segment.get("latex"), segment.get("mathml")
    if not isinstance(latex, str) or not latex.strip() or len(latex) > 12_000 \
            or not isinstance(raw, str) or len(raw) > MAX_MATHML:
        _fail(where, "公式缺少完整的本地排版结果，请重新打开组卷")
    if re.search(r"<!|<\?", raw):
        _fail(where, "公式 XML 不能含实体声明或处理指令")
    parser = etree.XMLParser(resolve_entities=False, load_dtd=False, no_network=True, huge_tree=False)
    try:
        root = etree.fromstring(raw.encode("utf-8"), parser)
        if root.tag != f"{{{MATH_NS}}}math":
            raise ValueError("root")
        nodes = list(root.iter())
        if len(nodes) > 3000 or any(not isinstance(node.tag, str) or not node.tag.startswith(f"{{{MATH_NS}}}")
                                   for node in nodes):
            raise ValueError("nodes")
        for node in nodes:
            if etree.QName(node).localname not in MATHML_ELEMENTS or len(list(node.iterancestors())) > 60 \
                    or any("href" in key.lower() or "src" in key.lower() or key.lower().startswith("on") for key in node.attrib):
                raise ValueError("attributes")
            # Upstream turns every menclose into a box, silently losing strikes.
            # A boxed formula is safe; cancellation/radical/arrow enclosures aren't.
            if word_math and etree.QName(node).localname == "menclose" and node.get("notation", "longdiv") != "box":
                raise ValueError("unsupported enclosure")
        annotations = root.findall(f".//{{{MATH_NS}}}annotation")
        if len(annotations) != 1 or annotations[0].get("encoding") != "application/x-tex" \
                or len(annotations[0]) or "".join(annotations[0].itertext()) != latex \
                or annotations[0].getparent().tag != f"{{{MATH_NS}}}semantics" or len(annotations[0].getparent()) < 2:
            raise ValueError("annotation")
        annotations[0].getparent().remove(annotations[0])
        if not word_math:
            # PDF renders canonical source again through KaTeX; validated metadata
            # stays in the segment, without imposing Word's OMML support limits.
            return None
        # Upstream SAX handler would print annotation characters as a side effect.
        omml = _convert_math_root(root)
        # Detect silent token loss in the converter, excluding intentional phantom.
        tokens = []
        for node in root.iter():
            if etree.QName(node).localname not in {"mi", "mn", "mo", "mtext", "ms"} \
                    or any(etree.QName(parent).localname == "mphantom" for parent in node.iterancestors()):
                continue
            token = "".join(node.itertext())
            parent = node.getparent()
            if parent is not None and parent.tag == f"{{{MATH_NS}}}mover" and parent.get("accent") == "true" \
                    and len(parent) == 2 and node is parent[1]:
                token = ACCENTS.get(token, token)
            tokens.append(token)
        output = "".join(omml.itertext()) + "".join(node.get(f"{{{OMML_NS}}}val", "") for node in omml.iter())
        for token in tokens:
            token = re.sub(r"[\s\u2061-\u2064]", "", token)
            if token and token not in re.sub(r"\s", "", output):
                raise ValueError("lost token")
        return omml
    except (ValueError, TypeError, RuntimeError, NotImplementedError, AssertionError, IndexError, KeyError,
            etree.XMLSyntaxError, xml.sax.SAXException):
        _fail(where, "这段公式暂不能准确转换为可编辑 Word 公式，请在预览核对并改用打印 / 保存 PDF" if word_math
              else "公式排版结果无效，请重新打开组卷预览后再导出 PDF")


def _segments(source, segments, start, end, where, offsets=None, *, word_math=True):
    offsets = offsets or _utf16_map(source, where)
    if not isinstance(segments, list) or len(segments) > 4000:
        _fail(where, "排版片段格式不正确")
    converted, cursor = [], start
    for segment in segments:
        if not isinstance(segment, dict) or set(segment) - SEGMENT_KEYS:
            _fail(where, "排版片段含不支持的字段")
        a, b = _bounds(segment, offsets, cursor, end, where)
        raw = source[offsets[a]:offsets[b]]
        kind = segment.get("type")
        if kind not in {"text", "math", "blank", "bracket", "parallelogram", "delimiter"}:
            _fail(where, "排版片段类型不正确")
        if "display" in segment and type(segment["display"]) is not bool:
            _fail(where, "公式展示方式不正确")
        if "displayGroup" in segment and (not isinstance(segment["displayGroup"], str) or len(segment["displayGroup"]) > 100):
            _fail(where, "公式分组不正确")
        if kind != "math" and ({"latex", "mathml"} & set(segment)):
            _fail(where, "普通文字不能携带替换公式")
        if kind == "delimiter" and raw not in {"$", "$$", "\\(", "\\)", "\\[", "\\]"}:
            _fail(where, "只能省略公式定界符，不能省略题目文字")
        if kind == "parallelogram" and not re.fullmatch(r"▱|\\(?:text|mathrm|mathbf)\s*\{\s*▱\s*\}", raw):
            _fail(where, "平行四边形符号与原文不符")
        if kind == "blank" and not re.fullmatch(r"(?:\\_){2,}|_{3,}", raw):
            _fail(where, "填空线与原文不符")
        if kind == "bracket" and not re.fullmatch(r"（[ \u3000]*）|\([ \u3000]+\)", raw):
            _fail(where, "答题括号与原文不符")
        if kind == "text" and re.search(r"\$[^$\n]+\$|\\[([]|\\[A-Za-z]{2,}", raw):
            _fail(where, "公式尚未转换，不能把 LaTeX 源码当作成功结果")
        if kind == "text" and re.search(r"<(?:table|tr|td|th)\b", raw, re.I):
            _fail(where, "表格尚未转换，不能把表格源码当作成功结果")
        converted.append({**segment, "raw": raw, "omml": _math(segment, where, word_math=word_math) if kind == "math" else None})
        cursor = b
    if cursor != end:
        _fail(where, "排版结果没有完整覆盖原文，请重新打开组卷")
    return converted


def _decoded(value):
    # Match QB's deliberately small entity decoder rather than reinterpret HTML.
    def decode(match):
        code = match.group(1).lower()
        named = {"amp": "&", "lt": "<", "gt": ">", "quot": '"', "apos": "'", "nbsp": " "}
        if code in named:
            return named[code]
        number = int(code[2:], 16) if code.startswith("#x") else int(code[1:])
        return chr(number) if 0 < number < 0x110000 else match.group(0)
    return re.sub(r"&(#x[0-9a-f]+|#\d+|amp|lt|gt|quot|apos|nbsp);", decode, value, flags=re.I)


def _table_source(raw, where):
    """Reconstruct authoritative cells; client rows cannot hide or alter cells."""
    if re.match(r"\s*<table\b", raw, re.I):
        if re.search(r"<(?:img|svg|math|script|iframe|object)\b", raw, re.I):
            _fail(where, "表格含嵌入图像或公式标记，暂不能保证完整 Word 导出")
        rows = []
        for row in re.finditer(r"<tr\b[^>]*>([\s\S]*?)</tr\s*>", raw, re.I):
            cells = []
            for cell in re.finditer(r"<(td|th)\b([^>]*)>([\s\S]*?)</\1\s*>", row[1], re.I):
                def span(name):
                    match = re.search(rf"\b{name}\s*=\s*[\"']?(\d{{1,2}})", cell[2], re.I)
                    return max(1, min(20, int(match[1]))) if match else 1
                value = re.sub(r"<br\s*/?>", " ", cell[3], flags=re.I)
                value = re.sub(r"<[^>]+>", "", value)
                cells.append({"source": re.sub(r"\s+", " ", _decoded(value)).strip(),
                              "header": cell[1].lower() == "th", "colspan": span("colspan"), "rowspan": span("rowspan")})
            if cells:
                rows.append(cells)
        if not rows or len(rows) > 60 or any(len(row) > 20 for row in rows):
            _fail(where, "表格结构不完整或超出 60 行、20 列，请分批整理")
        return rows
    lines = raw.split("\n")
    if len(lines) < 2 or any(not line.strip().startswith("|") for line in lines):
        _fail(where, "Markdown 表格结构不完整")
    rows = []
    for line in lines:
        line = line.strip()[1:]
        if line.endswith("|") and not line.endswith("\\|"):
            line = line[:-1]
        cells, start, in_math, index = [], 0, False, 0
        while index < len(line):
            char = line[index]
            if char == "\\" and index + 1 < len(line) and line[index + 1] == "|" and not in_math:
                index += 2
                continue
            if char == "$":
                in_math = not in_math
            if char == "|" and not in_math:
                cells.append(line[start:index].strip().replace("\\|", "|"))
                start = index + 1
            index += 1
        cells.append(line[start:].strip().replace("\\|", "|"))
        rows.append(cells)
    header = bool(rows[1] and any(rows[1]) and all(not cell or re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]))
    if header:
        rows = [rows[0], *rows[2:]]
    if len(rows) > 60 or any(len(row) > 20 for row in rows):
        _fail(where, "表格超出 60 行、20 列")
    return [[{"source": cell, "header": header and i == 0, "colspan": 1, "rowspan": 1} for cell in row]
            for i, row in enumerate(rows)]


def _field(value, official, where, *, word_math=True):
    if not isinstance(value, dict) or set(value) != {"source", "blocks"} or value.get("source") != official:
        _fail(where, "原文与当前入库版不一致，请重新打开组卷", 409)
    offsets = _utf16_map(official, where)
    end = max(offsets)
    blocks = value.get("blocks")
    if not isinstance(blocks, list) or len(blocks) > 1000:
        _fail(where, "排版块格式不正确")
    cursor, result = 0, []
    for block in blocks:
        if not isinstance(block, dict):
            _fail(where, "排版块格式不正确")
        a, b = _bounds(block, offsets, cursor, end, where)
        if block.get("type") == "text" and set(block) == {"type", "start", "end", "segments"}:
            result.append({"type": "text", "segments": _segments(official, block["segments"], a, b, where, offsets, word_math=word_math)})
        elif block.get("type") == "table" and set(block) == {"type", "start", "end", "rows"}:
            expected = _table_source(official[offsets[a]:offsets[b]], where)
            rows = block["rows"]
            if not isinstance(rows, list) or len(rows) != len(expected):
                _fail(where, "表格行数与原文不符")
            converted = []
            for given, actual in zip(rows, expected):
                if not isinstance(given, list) or len(given) != len(actual):
                    _fail(where, "表格格数与原文不符")
                cells = []
                for cell, reference in zip(given, actual):
                    if not isinstance(cell, dict) or set(cell) != {"source", "header", "colspan", "rowspan", "segments"} \
                            or any(cell.get(key) != reference[key] or type(cell.get(key)) is not type(reference[key]) for key in reference):
                        _fail(where, "表格内容或合并结构与原文不符")
                    cell_offsets = _utf16_map(cell["source"], where)
                    cells.append({**reference, "segments": _segments(cell["source"], cell["segments"], 0, max(cell_offsets), where, cell_offsets, word_math=word_math)})
                converted.append(cells)
            result.append({"type": "table", "rows": converted})
        else:
            _fail(where, "排版块含不支持的字段或类型")
        cursor = b
    if cursor != end:
        _fail(where, "排版结果没有完整覆盖原文，请重新打开组卷")
    return result


def _selected(content, extras, use_ai):
    if str(content.get("answer") or "").strip():
        return content, False
    ai = extras.get("ai_answer") if use_ai and isinstance(extras, dict) else None
    if isinstance(ai, dict) and (str(ai.get("answer") or "").strip() or str(ai.get("analysis") or "").strip()):
        return ai, True
    return (content, False) if str(content.get("analysis") or "").strip() else ({}, False)


def _images(publication, where):
    from PIL import Image
    images, total = [], 0
    figures = publication.content.get("figures", [])
    if not isinstance(figures, list):
        _fail(where, "配图快照不完整", 409)
    for index, figure in enumerate(figures):
        label = f"{where} · 配图 {index + 1}"
        if not isinstance(figure, dict) or figure.get("slot") not in {"stem", *"ABCDE", *(f"option_{key}" for key in "ABCDE")}:
            _fail(label, "配图位置无法识别，不能省略", 409)
        name = figure.get("file")
        if not isinstance(name, str) or not re.fullmatch(r"figure-\d{1,3}\.png", name):
            _fail(label, "缺少入库配图文件", 409)
        folder = Path(settings.DATA_ROOT) / "library" / str(publication.id)
        target = folder / name
        try:
            if target.is_symlink() or folder.is_symlink() or target.stat().st_size > MAX_IMAGE_BYTES:
                raise ValueError
            data = target.read_bytes()
            total += len(data)
            if len(data) > MAX_IMAGE_BYTES or total > MAX_TOTAL_IMAGES:
                raise ValueError
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(io.BytesIO(data)) as image:
                    size = image.size
                    if image.format != "PNG" or image.width < 1 or image.height < 1:
                        raise ValueError
                    image.verify()
            images.append({"slot": figure["slot"].removeprefix("option_"), "bytes": data, "size": size,
                           "name": name, "sha256": hashlib.sha256(data).hexdigest()})
        except (OSError, ValueError, Image.DecompressionBombError, Image.DecompressionBombWarning):
            _fail(label, "文件缺失、损坏或过大，请检查这道题的入库配图", 409)
    return images


def _capture(ids, rendered, options, output_format, *, word_math=None):
    word_math = output_format != "pdf" if word_math is None else word_math
    if not isinstance(rendered, dict) or set(rendered) != set(ids):
        raise ExportError("每道选题都需提供完整排版结果，请重新打开组卷")
    found = {str(item.id): item for item in PublishedQuestion.objects.filter(pk__in=ids)}
    use_ai = options["ai_answers"] and features.enabled("ai_answer")
    question_fields = output_format == "split" or options["document"] != "answers"
    answer_fields = output_format == "split" or options["document"] != "questions"
    captured, total_bytes = [], 0
    for position, key in enumerate(ids, 1):
        where = f"选题第 {position} 题"
        publication = found.get(key)
        if publication is None or publication.status != PublishedQuestion.Status.PUBLISHED:
            _fail(where, "这条入库版已撤回、被替代或不存在；请返回组卷处理后再导出", 409)
        if not isinstance(publication.content, dict) or not isinstance(publication.content.get("stem"), str):
            _fail(where, "题目快照不完整", 409)
        content, extras = deepcopy(publication.content), deepcopy(publication.extras)
        if any(content.get(name) is not None and not isinstance(content[name], str) for name in ("answer", "analysis", "origin")):
            _fail(where, "答案或题源快照格式不完整", 409)
        selected, is_ai = _selected(content, extras, use_ai)
        if any(selected.get(name) is not None and not isinstance(selected[name], str) for name in ("answer", "analysis")):
            _fail(where, "所选参考答案格式不完整", 409)
        if is_ai and ((selected.get("publication_id") and str(selected["publication_id"]) != key)
                      or (selected.get("fingerprint") and selected["fingerprint"] != library.generation_fingerprint(content, publication.id))):
            _fail(where, "AI 参考与这条入库版不匹配，请先核对或重新生成；也可关闭 AI 参考后导出", 409)
        required = {}
        if question_fields:
            required["stem"] = content["stem"]
            if not required["stem"].strip():
                _fail(where, "题干为空，不能导出空题", 409)
            option_values = content.get("options") or {}
            if not isinstance(option_values, dict) or set(option_values) - set("ABCDE"):
                _fail(where, "选项快照不完整", 409)
            for letter, text in option_values.items():
                if not isinstance(text, str):
                    _fail(where, "选项快照不完整", 409)
                if text.strip():
                    required[f"options.{letter}"] = text
            if options["origin"] and str(content.get("origin") or "").strip():
                required["origin"] = content["origin"]
        if answer_fields:
            required.update({name: str(selected.get(name) or "") for name in ("answer", "analysis")})
        given = rendered[key]
        if not isinstance(given, dict) or set(given) - FIELDS or not set(required) <= set(given):
            _fail(where, "缺少题干、选项或所选答案的排版字段")
        converted = {name: _field(given[name], text, f"{where} · {name}", word_math=word_math) for name, text in required.items()}
        # Optional empty fields sent by older clients must still match the official value.
        for name in set(given) - set(required):
            if name in {"answer", "analysis"}:
                official = str(selected.get(name) or "")
            elif name.startswith("options."):
                official = str((content.get("options") or {}).get(name[-1]) or "")
            else:
                official = str(content.get(name) or "")
            _field(given[name], official, f"{where} · {name}", word_math=word_math)
        images = _images(publication, where) if question_fields else []
        total_bytes += sum(len(image["bytes"]) for image in images)
        if total_bytes > MAX_TOTAL_IMAGES:
            _fail(where, "本次配图总量过大，请减少选题")
        captured.append({"id": key, "where": where, "content": content, "extras": extras, "selected": deepcopy(selected),
                         "ai": is_ai, "hash": publication.content_hash, "type": publication.question_type,
                         "fields": converted, "images": images})
    if (answer_fields and options["document"] == "answers") or output_format == "split":
        if not any(str(item["selected"].get("answer") or "").strip() or str(item["selected"].get("analysis") or "").strip() for item in captured):
            raise ExportError("这些题没有原卷答案或所选 AI 参考，请改选题目卷，或先补齐答案再分卷导出")
    return captured, use_ai


def _recheck(captured, options, use_ai):
    found = {str(item.id): item for item in PublishedQuestion.objects.filter(pk__in=[item["id"] for item in captured])}
    current_ai = options["ai_answers"] and features.enabled("ai_answer")
    for item in captured:
        live = found.get(item["id"])
        if live is None or live.status != PublishedQuestion.Status.PUBLISHED or live.content_hash != item["hash"] \
                or live.content != item["content"] or live.question_type != item["type"]:
            _fail(item["where"], "导出期间入库版本发生变化，请重新打开组卷", 409)
        selected, is_ai = _selected(live.content, live.extras, current_ai)
        if selected != item["selected"] or is_ai != item["ai"] or use_ai != current_ai:
            _fail(item["where"], "导出期间答案或 AI 设置发生变化，请重新打开组卷", 409)
        for before, after in zip(item["images"], _images(live, item["where"]) if item["images"] else []):
            if before["sha256"] != after["sha256"]:
                _fail(item["where"], "导出期间配图发生变化，请重新打开组卷", 409)


def _font(run, size, bold=False):
    from docx.shared import Pt
    from docx.oxml.ns import qn
    run.font.name = "Cambria"
    run.font.size = Pt(size)
    run.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    return run


def _write_segments(paragraph, segments, size, bold=False):
    for segment in segments:
        kind = segment["type"]
        if kind == "math":
            # Native oMath, not screenshots. Display grouping is kept in one paragraph.
            paragraph._p.append(deepcopy(segment["omml"]))
        elif kind != "delimiter":
            shown = "________" if kind == "blank" else "（\u3000\u3000）" if kind == "bracket" else "▱" if kind == "parallelogram" else segment["raw"]
            _font(paragraph.add_run(shown), size, bold)


def _write_table(document, rows, size, where):
    from docx.shared import Mm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    # Place rowspan cells into a complete grid; overlapping/out-of-bounds spans abort.
    grid, placements, width = {}, [], 0
    for r, row in enumerate(rows):
        column = 0
        for cell in row:
            while (r, column) in grid:
                column += 1
            h, w = cell["rowspan"], cell["colspan"]
            if r + h > len(rows) or column + w > 40:
                _fail(where, "表格合并单元格超出边界")
            for y in range(r, r + h):
                for x in range(column, column + w):
                    if (y, x) in grid:
                        _fail(where, "表格合并单元格重叠")
                    grid[y, x] = True
            placements.append((r, column, cell))
            column += w
            width = max(width, column)
    if width < 1:
        _fail(where, "表格没有可导出的单元格")
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.autofit = False
    for column in table.columns:
        column.width = Mm(178 / width)
    for r, c, cell in placements:
        target = table.cell(r, c)
        if cell["rowspan"] > 1 or cell["colspan"] > 1:
            target = target.merge(table.cell(r + cell["rowspan"] - 1, c + cell["colspan"] - 1))
        target.text = ""
        _write_segments(target.paragraphs[0], cell["segments"], size, cell["header"])
        if cell["header"]:
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "F1F5F3")
            target._tc.get_or_add_tcPr().append(shade)
    for row in table.rows:
        avoid = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(avoid)


def _write_field(document, blocks, size, prefix="", where="", lead_segments=None):
    from docx.shared import Pt
    paragraph = None
    if prefix:
        paragraph = document.add_paragraph()
        _font(paragraph.add_run(prefix), size)
    if lead_segments:
        paragraph = paragraph or document.add_paragraph()
        _write_segments(paragraph, lead_segments, size)
    for block in blocks:
        if block["type"] == "table":
            _write_table(document, block["rows"], size, where)
            paragraph = None
        else:
            display_group = None
            for segment in block["segments"]:
                if segment["type"] == "delimiter":
                    continue
                group = segment.get("displayGroup") or ("display" if segment.get("display") else None)
                if group != display_group and (group is not None or display_group is not None):
                    paragraph = None
                # Source newlines become real Word paragraphs so compact mode
                # can break between subquestions, rather than keeping a huge
                # paragraph with many manual line breaks on one page. Empty
                # source lines use the normal paragraph gap, not extra full-
                # height empty paragraphs; inline spaces still remain intact.
                parts = re.split(r"\r\n|[\r\n]", segment["raw"]) if segment["type"] == "text" else [None]
                for index, part in enumerate(parts):
                    if part is None or (part and (paragraph is not None or part.strip())):
                        paragraph = paragraph or document.add_paragraph()
                        if group is not None:
                            paragraph.alignment = 1
                        _write_segments(paragraph, [{**segment, "raw": part} if part is not None else segment], size)
                        paragraph.paragraph_format.space_after = Pt(PARAGRAPH_GAP_PT)
                        paragraph.paragraph_format.line_spacing = LINE_SPACING
                    if index < len(parts) - 1:
                        paragraph = None
                display_group = group
            paragraph = None


def _write_images(document, item, slot):
    from docx.shared import Mm
    for image in item["images"]:
        if image["slot"] != slot:
            continue
        width, height = image["size"]
        # 150 dpi natural size, capped to the A4 text area and one-page height.
        scale = min(178 / width, 210 / height, 25.4 / 150)
        paragraph = document.add_paragraph()
        paragraph.add_run().add_picture(io.BytesIO(image["bytes"]), width=Mm(width * scale), height=Mm(height * scale))
        paragraph.paragraph_format.keep_together = True


def _text_em(value):
    """Conservative Cambria/Chinese width in em, independent of LaTeX spelling."""
    total = 0.0
    for char in value:
        if unicodedata.combining(char):
            continue
        total += 1.0 if unicodedata.east_asian_width(char) in {"W", "F"} else 0.30 if char.isspace() else 0.60
    return total


def _math_extent(element):
    """Approximate native formula width/height for safe column selection.

    Fractions, radicals, scripts and matrices are measured structurally; raw
    command length is a poor proxy. This errs toward fewer columns and never
    changes font size or mathematical content.
    """
    from lxml import etree
    name = etree.QName(element).localname
    if name.endswith("Pr"):
        return 0.0, 0.0
    if name == "t":
        return _text_em(element.text or ""), 1.0
    children = {etree.QName(child).localname: child for child in element}
    def size(key):
        return _math_extent(children[key]) if key in children else (0.0, 0.0)
    if name == "f":
        numerator, denominator = size("num"), size("den")
        return max(numerator[0], denominator[0]) + 0.4, numerator[1] + denominator[1] + 0.3
    if name == "rad":
        base, degree = size("e"), size("deg")
        return base[0] + 1.0 + degree[0] * 0.35, max(1.3, base[1] + 0.25)
    if name in {"sSup", "sSub", "sSubSup"}:
        base, upper, lower = size("e"), size("sup"), size("sub")
        return base[0] + 0.7 * max(upper[0], lower[0]), base[1] + 0.5 * (upper[1] + lower[1])
    if name in {"acc", "bar", "groupChr"}:
        base = size("e")
        return base[0], base[1] + 0.4
    if name == "d":
        base = size("e")
        return base[0] + 1.0, base[1]
    if name in {"limLow", "limUpp"}:
        base, limit = size("e"), size("lim")
        return max(base[0], limit[0] * 0.75), base[1] + limit[1] * 0.6
    if name == "nary":
        base, upper, lower = size("e"), size("sup"), size("sub")
        return 1.0 + max(upper[0], lower[0]) * 0.65 + base[0], max(base[1], 1.5 + (upper[1] + lower[1]) * 0.6)
    if name in {"m", "eqArr"}:
        rows = [child for child in element if etree.QName(child).localname in {"mr", "e"}]
        metrics = [_math_extent(row) for row in rows]
        return max((metric[0] for metric in metrics), default=1.0), sum(metric[1] + 0.2 for metric in metrics)
    metrics = [_math_extent(child) for child in element]
    return sum(metric[0] for metric in metrics), max((metric[1] for metric in metrics), default=0.0)


def _option_columns(item, options, *, text_width_mm=178):
    values = item["content"].get("options") or {}
    letters = [key for key in "ABCDE" if str(values.get(key) or "").strip()]
    if len(letters) < 2 or any(image["slot"] != "stem" for image in item["images"]):
        return 1
    widths = []
    for letter in letters:
        blocks = item["fields"].get(f"options.{letter}", [])
        if re.search(r"[\r\n]", values[letter]) or len(blocks) != 1 or blocks[0]["type"] != "text":
            return 1
        width = 1.8  # letter, period and breathing room
        for segment in blocks[0]["segments"]:
            if segment.get("display") or segment.get("displayGroup"):
                return 1
            if segment["type"] == "math":
                if segment["omml"] is None:
                    return 1
                extent = _math_extent(segment["omml"])
                if extent[1] > 2.6:
                    return 1
                width += extent[0]
            elif segment["type"] != "delimiter":
                width += _text_em(segment["raw"]) if segment["type"] == "text" else 4.5
        widths.append(width * options["font_size"] * 25.4 / 72)
    requested = options.get("option_overrides", {}).get(item["id"], options.get("option_layout", "auto"))
    candidates = (4, 2) if requested != "two" and len(letters) == 4 else (2,)
    for columns in candidates:
        # Cell margins and a 10% reserve absorb font/rendering variation.
        if max(widths) <= (text_width_mm / columns - 4.0) * 0.90:
            return columns
    return 1


def _write_compact_options(document, item, size, columns, *, text_width_mm=178):
    from docx.shared import Mm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    letters = [key for key in "ABCDE" if f"options.{key}" in item["fields"]]
    table = document.add_table(rows=(len(letters) + columns - 1) // columns, cols=columns)
    table.autofit = False
    for column in table.columns:
        column.width = Mm(text_width_mm / columns)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + side)
        border.set(qn("w:val"), "nil")
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for index, letter in enumerate(letters):
        paragraph = table.cell(index // columns, index % columns).paragraphs[0]
        _font(paragraph.add_run(letter + ". "), size)
        _write_segments(paragraph, item["fields"][f"options.{letter}"][0]["segments"], size)
        paragraph.paragraph_format.line_spacing = LINE_SPACING
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _ordered(captured):
    known = {key for key, _ in GROUPS}
    result = [(name, [item for item in captured if item["type"] == key]) for key, name in GROUPS]
    result.append(("其他", [item for item in captured if item["type"] not in known]))
    return [(name, group) for name, group in result if group]


def _keep_unit(document, start):
    """Keep a question together when it fits; Word may split an over-page item.

    Cell paragraphs also carry keepNext so a short table and its following
    diagram/options stay with the stem. The final paragraph releases the chain.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    elements = list(document._element.body)[start:-1]  # last child is sectPr
    paragraphs = []
    for element in elements:
        if element.tag == qn("w:p"):
            paragraphs.append(element)
        elif element.tag == qn("w:tbl"):
            paragraphs.extend(element.findall(".//" + qn("w:p")))
    for index, paragraph in enumerate(paragraphs):
        props = paragraph.get_or_add_pPr()
        for name in ("w:keepNext", "w:keepLines"):
            existing = props.find(qn(name))
            if existing is not None:
                props.remove(existing)
            node = OxmlElement(name)
            node.set(qn("w:val"), "0" if name == "w:keepNext" and index == len(paragraphs) - 1 else "1")
            props.append(node)
    # A final table releases every cell in its final physical row, otherwise a
    # neighbouring cell's keepNext can accidentally chain the following question.
    if elements and elements[-1].tag == qn("w:tbl"):
        rows = elements[-1].findall(qn("w:tr"))
        if rows:
            for paragraph in rows[-1].findall(".//" + qn("w:p")):
                keep = paragraph.get_or_add_pPr().find(qn("w:keepNext"))
                if keep is not None:
                    keep.set(qn("w:val"), "0")


def _paragraph_measure(paragraph, size, width_mm=178):
    from docx.oxml.ns import qn
    import math
    spacing = paragraph.find(qn("w:pPr") + "/" + qn("w:spacing"))
    gap = int(spacing.get(qn("w:after"), str(PARAGRAPH_GAP_PT * 20))) * 25.4 / 1440 if spacing is not None \
        else PARAGRAPH_GAP_PT * 25.4 / 72
    drawings = paragraph.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent")
    if drawings:
        return max(int(node.get("cy", "0")) / 36_000 for node in drawings) + gap
    text = "".join(node.text or "" for node in paragraph.findall(".//" + qn("w:t")))
    metrics = [_math_extent(node) for node in paragraph.findall(".//" + qn("m:oMath"))]
    width = (_text_em(text) + sum(metric[0] for metric in metrics)) * size * 25.4 / 72
    height = max(LINE_SPACING, max((metric[1] for metric in metrics), default=1.0)) * size * 25.4 / 72
    return max(1, math.ceil(width / width_mm)) * height + gap


def _unit_height(elements, size):
    from docx.oxml.ns import qn
    total = 0.0
    for element in elements:
        if element.tag == qn("w:p"):
            total += _paragraph_measure(element, size)
        elif element.tag == qn("w:tbl"):
            columns = max(1, len(element.findall(qn("w:tblGrid") + "/" + qn("w:gridCol"))))
            for row in element.findall(qn("w:tr")):
                total += max((sum(_paragraph_measure(paragraph, size, 178 / columns - 4)
                                  for paragraph in cell.findall(qn("w:p")))
                              for cell in row.findall(qn("w:tc"))), default=0)
    return total


def _pagination(document, start, mode, size):
    """Compact flows at paragraph/row boundaries; short keep-mode items bind."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    elements = list(document._element.body)[start:-1]
    if mode == "keep" and _unit_height(elements, size) <= 230:
        _keep_unit(document, start)
        return
    for element in elements:
        paragraphs = [element] if element.tag == qn("w:p") else element.findall(".//" + qn("w:p"))
        for paragraph in paragraphs:
            props = paragraph.get_or_add_pPr()
            plain = "".join(node.text or "" for node in paragraph.findall(".//" + qn("w:t"))).strip()
            atomic = element.tag == qn("w:tbl") or bool(paragraph.findall(".//" + qn("w:drawing"))) \
                or (bool(paragraph.findall(".//" + qn("m:oMath"))) and not plain.strip("▱"))
            for name, value in (("w:keepNext", False), ("w:keepLines", atomic), ("w:widowControl", True)):
                existing = props.find(qn(name))
                if existing is not None:
                    props.remove(existing)
                node = OxmlElement(name)
                node.set(qn("w:val"), "1" if value else "0")
                props.append(node)
    # An isolated number/origin lead must stay with the start of its actual stem.
    # Only a one-line caption stays with the immediately following diagram/table.
    # A normal multi-line stem must remain breakable in compact mode; plain-text
    # length alone misses the width of its native formulas.
    for index, element in enumerate(elements[:-1]):
        if element.tag != qn("w:p"):
            continue
        plain = "".join(node.text or "" for node in element.findall(".//" + qn("w:t"))).strip()
        next_element = elements[index + 1]
        lead_only = bool(re.fullmatch(r"\d+\.(?:（AI 参考 · 未核对）)?\s*", plain))
        next_visual = next_element.tag == qn("w:tbl") or bool(next_element.findall(".//" + qn("w:drawing")))
        formula_width = sum(_math_extent(node)[0] for node in element.findall(".//" + qn("m:oMath")))
        width_mm = (_text_em(plain) + formula_width) * size * 25.4 / 72
        short_caption = bool(plain and len(plain) <= 40 and width_mm <= 178)
        if lead_only or (short_caption and next_visual):
            element.get_or_add_pPr().find(qn("w:keepNext")).set(qn("w:val"), "1")


def _question_gap(document):
    from docx.shared import Pt
    from docx.oxml.ns import qn
    last = list(document._element.body)[-2]
    if last.tag == qn("w:tbl"):
        # Word requires a paragraph after a table; use its invisible line for the
        # same external question gap as a normal paragraph, without resizing text.
        if "QB Question Gap" not in document.styles:
            from docx.enum.style import WD_STYLE_TYPE
            style = document.styles.add_style("QB Question Gap", WD_STYLE_TYPE.PARAGRAPH)
            style.font.size = Pt(1)
            style.paragraph_format.line_spacing = Pt(1)
            style.paragraph_format.space_after = Pt(QUESTION_GAP_PT)
        document.add_paragraph(style="QB Question Gap")
    else:
        from docx.text.paragraph import Paragraph
        paragraph = Paragraph(last, document)
        after = paragraph.paragraph_format.space_after
        if after is None or after < Pt(QUESTION_GAP_PT):
            paragraph.paragraph_format.space_after = Pt(QUESTION_GAP_PT)


def _question_page_break(document, start):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    elements = list(document._element.body)[start:-1]
    first = next((element for element in elements if element.tag == qn("w:p")), None)
    if first is not None:
        first.get_or_add_pPr().append(OxmlElement("w:pageBreakBefore"))


def _document(captured, title, options, mode):
    from docx import Document
    from docx.shared import Mm, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(16)
    section.header_distance = section.footer_distance = Mm(8)
    size = options["font_size"]
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Cambria", Pt(size)
    normal.paragraph_format.line_spacing = LINE_SPACING
    normal.paragraph_format.space_after = Pt(PARAGRAPH_GAP_PT)
    normal.paragraph_format.widow_control = True
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    # Word math defaults, including editable equations, follow the selected size.
    default = document.styles["Default Paragraph Font"]
    default.font.name, default.font.size = "Cambria", Pt(size)
    heading = document.add_paragraph()
    heading.alignment = 1
    heading.paragraph_format.keep_with_next = True
    _font(heading.add_run(title + (" · 答案解析" if mode == "answers" else "")), size + 6, True)
    if options["student_info"] and mode != "answers":
        info = document.add_paragraph()
        info.paragraph_format.keep_with_next = True
        _font(info.add_run("姓名：____________    班级：____________    得分：____________"), size)
    numbered, number = [], 0
    for index, (name, group) in enumerate(_ordered(captured)):
        section_title = None
        if mode != "answers":
            section_title = document.add_paragraph()
            section_title.paragraph_format.keep_with_next = True
            _font(section_title.add_run(f"{'一二三四五六'[index]}、{name}"), size + 1, True)
        for item in group:
            number += 1
            numbered.append((number, item))
            if mode == "answers":
                continue
            unit_start = len(document._element.body) - 1
            prefix = f"{number}. "
            lead_segments = None
            if "origin" in item["fields"]:
                origin_blocks = item["fields"]["origin"]
                if len(origin_blocks) == 1 and origin_blocks[0]["type"] == "text":
                    lead_segments = [{"type": "text", "raw": "（"}, *origin_blocks[0]["segments"], {"type": "text", "raw": "）"}]
                else:
                    _write_field(document, origin_blocks, size, prefix=prefix + "（题源）", where=item["where"])
                    prefix = ""
            _write_field(document, item["fields"]["stem"], size, prefix=prefix, where=item["where"], lead_segments=lead_segments)
            _write_images(document, item, "stem")
            columns = _option_columns(item, options)
            if columns > 1:
                _write_compact_options(document, item, size, columns)
            else:
                for letter in "ABCDE":
                    if f"options.{letter}" in item["fields"]:
                        _write_field(document, item["fields"][f"options.{letter}"], size, prefix=f"{letter}. ", where=item["where"])
                    elif any(image["slot"] == letter for image in item["images"]):
                        _font(document.add_paragraph().add_run(f"{letter}."), size)
                    _write_images(document, item, letter)
            if item["type"] == "free_response" and options["answer_space"] != "none":
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Mm(30 if options["answer_space"] == "medium" else 60)
            _question_gap(document)
            _pagination(document, unit_start, options["pagination"], size)
            if number > 1 and item["id"] in options["question_breaks"]:
                if item is group[0] and section_title is not None:
                    section_title.paragraph_format.page_break_before = True
                else:
                    _question_page_break(document, unit_start)
    if mode != "questions":
        if mode == "combined":
            document.add_page_break()
        answer_heading = document.add_paragraph()
        answer_heading.paragraph_format.keep_with_next = True
        _font(answer_heading.add_run("参考答案与解析"), size + 1, True)
        for number, item in numbered:
            unit_start = len(document._element.body) - 1
            label = f"{number}." + ("（AI 参考 · 未核对）" if item["ai"] else "")
            selected = item["selected"]
            if not str(selected.get("answer") or "").strip() and not str(selected.get("analysis") or "").strip():
                _font(document.add_paragraph().add_run(label + "（原卷未提供答案）"), size)
                _question_gap(document)
                _pagination(document, unit_start, options["pagination"], size)
                continue
            _write_field(document, item["fields"]["answer"], size, prefix=label + " ", where=item["where"])
            if str(selected.get("analysis") or "").strip():
                _write_field(document, item["fields"]["analysis"], size, prefix="解析：", where=item["where"])
            _question_gap(document)
            _pagination(document, unit_start, options["pagination"], size)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    _font(footer.add_run(f"共 {len(captured)} 题 · 第 "), 9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    _font(footer.add_run(" 页"), 9)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def safe_filename(title):
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", title).rstrip(" .") or "练习"
    if re.match(r"^(?:CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\.|$)", value, re.I):
        value = "练习-" + value
    return value[:100]


def export(payload):
    if not isinstance(payload, dict) or set(payload) - {"ids", "title", "print_options", "rendered_fields", "format"}:
        raise ExportError("导出请求包含不支持的字段")
    try:
        ids = normalize_ids(payload.get("ids"))
        title = _title(payload.get("title", "练习"))
        options = _print_options(payload.get("print_options", {}), ids=ids)
    except (BrowseError, DraftError) as error:
        raise ExportError(str(error)) from None
    if not ids:
        raise ExportError("请先把题目加入题篮")
    _utf16_map(title, "卷名")
    if len(ids) != len(payload["ids"]):
        raise ExportError("选题编号重复，请重新打开组卷")
    output_format = payload.get("format", "docx")
    if not isinstance(output_format, str) or output_format not in {"docx", "split"}:
        raise ExportError("导出格式只能选 docx 或 split")
    captured, use_ai = _capture(ids, payload.get("rendered_fields"), options, output_format)
    name = safe_filename(title)
    if output_format == "split":
        questions = _document(captured, title, options, "questions")
        answers = _document(captured, title, options, "answers")
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(name + "-题目卷.docx", questions)
            archive.writestr(name + "-答案解析卷.docx", answers)
        data, filename, mime = stream.getvalue(), name + "-分卷.zip", "application/zip"
    else:
        data = _document(captured, title, options, options["document"])
        filename, mime = name + ".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    _recheck(captured, options, use_ai)
    return data, filename, mime, len(ids)


@csrf_exempt
def export_docx_view(request):
    if request.method != "POST":
        return HttpResponseNotAllowed(["POST"])
    from .views import _body, _guard
    rejected = _guard(request)
    if rejected:
        return rejected
    if request.META.get("REMOTE_ADDR") not in {"127.0.0.1", "::1"}:
        return JsonResponse({"error": "导出需从本机题库页面发起"}, status=403)
    try:
        payload = _body(request, limit=MAX_BODY)
        if payload is None:
            raise ExportError("导出内容格式不正确或超过 2 MiB，请减少选题、分批导出", 413 if len(request.body) > MAX_BODY else 400)
        data, filename, mime, count = export(payload)
    except RequestDataTooBig:
        return JsonResponse({"error": "导出内容超过 2 MiB，请减少选题、分批导出"}, status=413)
    except ExportError as error:
        return JsonResponse({"error": str(error)}, status=error.status)
    response = HttpResponse(data, content_type=mime)
    response["Content-Disposition"] = "attachment; filename=practice." + ("zip" if mime == "application/zip" else "docx") + "; filename*=UTF-8''" + quote(filename)
    response["X-Question-Count"] = str(count)
    response["Cache-Control"] = "no-store"
    return response
